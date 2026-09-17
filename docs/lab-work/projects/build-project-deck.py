#!/usr/bin/env python3
"""
build-project-deck.py — turn a project report markdown file into a stakeholder deck.

Usage:
    python3 build-project-deck.py project-cindy-pawford.md [-o output.pptx]

Reads a markdown file written against PROJECT-TEMPLATE.md and emits a PowerPoint
deck styled to match "Autonomous Publishing Pattern v1.pptx" and
"Agent SDLC Process v1.1.pptx".

What it reads:
  * YAML-ish front matter: project, status, version, owner, started, last_updated,
    appliance, related_decks
  * "## 1. One-liner"            -> title slide subtitle
  * "## 2. Why this exists"      -> numbered research-question cards
  * "## 3. Scope"                -> two-column table
  * "## 4. Architecture"         -> table + "The loop" callout
  * "## 5. Related patterns"     -> table
  * "## 6. Current state"        -> bold-led blocks on the status slide
  * "## 7. Milestones"           -> table
  * "## 8. What I've learned"    -> table
  * "## 9. Open questions"       -> table
  * "## 10. Daily log"           -> one slide per recent entry (default: 3)

Section numbers matter; section titles do not. Requires python-pptx.
"""

import argparse
import re
import sys
from datetime import date
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

# ----------------------------------------------------------------------------
# Design tokens — lifted from the existing Patterns decks
# ----------------------------------------------------------------------------
DARK = "111827"
INK = "1F2937"
BODY = "374151"
MUTED = "6B7280"
BORDER = "E5E7EB"
PANEL = "F3F4F6"
WHITE = "FFFFFF"
ON_DARK_SUB = "93C5FD"
ON_DARK_MUTED = "9CA3AF"

AMBER = "B45309"
TEAL = "0E7490"
VIOLET = "6D28D9"
ACCENTS = [VIOLET, TEAL, AMBER]

HEAD_FONT = "Cambria"
BODY_FONT = "Calibri"

M = 0.62                # left margin
CW = 12.09              # content width
SLIDE_W, SLIDE_H = 13.3333, 7.5

GREEN = "15803D"

# Emoji status markers in the markdown are rendered as coloured words, because
# emoji glyph coverage in PowerPoint is not dependable across platforms.
STATUS_WORD = {
    "✅": ("Done", GREEN),
    "🟡": ("In flight", AMBER),
    "🔴": ("Planned", MUTED),
    "⬜": ("Planned", MUTED),
}


def rgb(h):
    return RGBColor.from_string(h)


# ----------------------------------------------------------------------------
# Markdown parsing
# ----------------------------------------------------------------------------
def parse_front_matter(text):
    meta, decks = {}, []
    m = re.match(r"^---\n(.*?)\n---\n", text, re.S)
    if not m:
        return meta, decks, text
    block = m.group(1)
    cur = None
    for line in block.splitlines():
        if not line.strip():
            continue
        if re.match(r"^\s+-\s", line) and cur == "related_decks":
            decks.append({})
            kv = line.strip()[2:]
            if ":" in kv:
                k, v = kv.split(":", 1)
                decks[-1][k.strip()] = v.strip()
        elif re.match(r"^\s+\w+:", line) and cur == "related_decks" and decks:
            k, v = line.strip().split(":", 1)
            decks[-1][k.strip()] = v.strip()
        elif re.match(r"^\w[\w_]*:", line):
            k, v = line.split(":", 1)
            cur = k.strip()
            v = v.strip()
            if v:
                meta[cur] = v.strip("[]")
    return meta, decks, text[m.end():]


def split_sections(body):
    """Return {section_number: {'title': str, 'body': str}}."""
    out = {}
    parts = re.split(r"^##\s+(\d+)\.\s*(.+)$", body, flags=re.M)
    for i in range(1, len(parts) - 1, 3):
        out[int(parts[i])] = {"title": parts[i + 1].strip(), "body": parts[i + 2]}
    return out


def strip_md(s):
    # protect code spans so literal asterisks inside them survive
    spans = []

    def stash(m):
        spans.append(m.group(1))
        return "\x00%d\x00" % (len(spans) - 1)

    s = re.sub(r"`([^`]+)`", stash, s)
    s = re.sub(r"!?\[([^\]]*)\]\([^)]*\)", r"\1", s)
    s = re.sub(r"\*\*(.+?)\*\*", r"\1", s)
    s = re.sub(r"(?<!\*)\*(?!\*)(.+?)(?<!\*)\*(?!\*)", r"\1", s)
    s = re.sub(r"\x00(\d+)\x00", lambda m: spans[int(m.group(1))], s)
    return s.strip()


def unwrap(s):
    """Join hard-wrapped lines inside a paragraph."""
    return re.sub(r"\s*\n\s*", " ", s).strip()


def parse_tables(section_body):
    """Return a list of tables; each is a list of rows of cell strings."""
    tables, cur = [], []
    for line in section_body.splitlines():
        ls = line.strip()
        if ls.startswith("|") and ls.endswith("|"):
            cells = [c.strip() for c in ls.strip("|").split("|")]
            if all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                continue
            cur.append([strip_md(c) for c in cells])
        elif cur:
            tables.append(cur)
            cur = []
    if cur:
        tables.append(cur)
    return tables


def parse_numbered(section_body):
    """Return [(n, text)] for '1. text' items, supporting wrapped lines."""
    items = []
    for m in re.finditer(r"^(\d+)\.\s+(.*?)(?=\n\s*\n|\n\d+\.\s|\Z)", section_body, re.S | re.M):
        items.append((m.group(1), unwrap(strip_md(m.group(2)))))
    return items


def parse_bold_blocks(section_body):
    """Return [(label, text)] for paragraphs shaped '**Label.** body'."""
    blocks = []
    for m in re.finditer(r"^\*\*(.+?)\.?\*\*\s*(.*?)(?=\n\s*\n|\Z)", section_body, re.S | re.M):
        label = m.group(1).strip().rstrip(".")
        blocks.append((label, unwrap(strip_md(m.group(2)))))
    return blocks


def parse_log(section_body, limit=3):
    """Return [{'date','headline','fields':[(label,text)]}] newest-first."""
    entries = []
    chunks = re.split(r"^###\s+(.+)$", section_body, flags=re.M)
    for i in range(1, len(chunks) - 1, 2):
        head = chunks[i].strip()
        m = re.match(r"([\d]{4}-[\d]{2}-[\d]{2})\s*[—–-]\s*(.*)", head)
        d, headline = (m.group(1), m.group(2)) if m else ("", head)
        entries.append({"date": d, "headline": strip_md(headline),
                        "fields": parse_bold_blocks(chunks[i + 1])})
    return entries[:limit]


# ----------------------------------------------------------------------------
# Drawing primitives
# ----------------------------------------------------------------------------
class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width = Inches(SLIDE_W)
        self.prs.slide_height = Inches(SLIDE_H)
        self.blank = self.prs.slide_layouts[6]

    def slide(self, dark=False):
        s = self.prs.slides.add_slide(self.blank)
        if dark:
            self._bg(s, DARK)
        return s

    @staticmethod
    def _bg(slide, color):
        bg = slide.background
        fill = bg.fill
        fill.solid()
        fill.fore_color.rgb = rgb(color)

    # -- text -----------------------------------------------------------
    @staticmethod
    def text(slide, l, t, w, h, runs, size=13, bold=False, color=BODY,
             font=BODY_FONT, align=None, spc=None, line=None, anchor=MSO_ANCHOR.TOP):
        box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = anchor
        paras = runs if isinstance(runs, list) else [runs]
        for i, ptext in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            if align is not None:
                p.alignment = align
            if line:
                p.line_spacing = Pt(line)
            r = p.add_run()
            r.text = ptext
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.name = font
            r.font.color.rgb = rgb(color)
            if spc:
                r.font._rPr.set("spc", str(int(spc)))
        return box

    # -- shapes ---------------------------------------------------------
    @staticmethod
    def rect(slide, l, t, w, h, fill=WHITE, line=None):
        sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l), Inches(t),
                                    Inches(w), Inches(h))
        sh.shadow.inherit = False
        if fill is None:
            sh.fill.background()
        else:
            sh.fill.solid()
            sh.fill.fore_color.rgb = rgb(fill)
        if line:
            sh.line.color.rgb = rgb(line)
            sh.line.width = Pt(1)
        else:
            sh.line.fill.background()
        return sh

    def swatches(self, slide, t=5.15):
        for i, c in enumerate([AMBER, TEAL, VIOLET]):
            self.rect(slide, M + i * 0.28, t, 0.15, 0.15, fill=c)

    # -- composites -----------------------------------------------------
    def heading(self, slide, title, kicker=None):
        self.text(slide, M, 0.38, CW, 0.62, title, size=30, bold=True,
                  color=INK, font=HEAD_FONT, anchor=MSO_ANCHOR.MIDDLE)
        if kicker:
            self.text(slide, M, 1.00, CW, 0.36, kicker, size=14, color=MUTED,
                      anchor=MSO_ANCHOR.MIDDLE)

    def divider(self, part, title, sub):
        s = self.slide(dark=True)
        self.text(s, M, 2.55, 8.0, 0.35, part.upper(), size=14, bold=True,
                  color=AMBER, spc=300, anchor=MSO_ANCHOR.MIDDLE)
        self.text(s, M, 3.05, 11.5, 0.95, title, size=40, bold=True, color=WHITE,
                  font=HEAD_FONT, anchor=MSO_ANCHOR.MIDDLE)
        self.text(s, M, 4.15, 10.5, 0.60, sub, size=16, color=ON_DARK_SUB,
                  anchor=MSO_ANCHOR.MIDDLE)
        self.swatches(s)
        return s

    def numbered_card(self, slide, idx, t, h, title, body, accent):
        self.rect(slide, M, t, CW, h, fill=WHITE, line=BORDER)
        self.rect(slide, M, t, 0.82, h, fill=accent)
        self.text(slide, M, t + (h - 0.7) / 2, 0.82, 0.7, str(idx), size=26, bold=True,
                  color=WHITE, font=HEAD_FONT, align=PP_ALIGN.CENTER,
                  anchor=MSO_ANCHOR.MIDDLE)
        if title:
            self.text(slide, 1.76, t + 0.18, CW - 1.30, 0.42, title, size=17.5,
                      bold=True, color=INK, font=HEAD_FONT, anchor=MSO_ANCHOR.MIDDLE)
            self.text(slide, 1.76, t + 0.62, CW - 1.30, h - 0.78, body, size=12.5,
                      color=BODY, line=18)
        else:
            size = 12.5 if len(body) <= 230 else 11.5
            self.text(slide, 1.76, t + 0.14, CW - 1.30, h - 0.26, body, size=size,
                      color=BODY, line=17, anchor=MSO_ANCHOR.MIDDLE)

    def callout(self, slide, t, title, body, h=1.05, fill=PANEL, size=12):
        self.rect(slide, M, t, CW, h, fill=fill)
        self.text(slide, M + 0.28, t + 0.15, CW - 0.56, 0.28, title, size=13.5,
                  bold=True, color=INK, font=HEAD_FONT)
        self.text(slide, M + 0.28, t + 0.45, CW - 0.56, h - 0.56, body, size=size,
                  color=BODY, line=16)

    def table(self, slide, t, rows, col_w=None, max_h=None, body_size=11.5):
        ncols = max(len(r) for r in rows)
        rows = [r + [""] * (ncols - len(r)) for r in rows]
        if not col_w:
            col_w = [CW / ncols] * ncols
        scale = CW / sum(col_w)
        col_w = [c * scale for c in col_w]

        nrows = len(rows)
        row_h = 0.40
        if max_h:
            row_h = min(0.46, max(0.28, max_h / nrows))
        gf = slide.shapes.add_table(nrows, ncols, Inches(M), Inches(t),
                                    Inches(CW), Inches(row_h * nrows))
        tbl = gf.table
        # kill banding / built-in style emphasis
        tblPr = tbl._tbl.find(qn("a:tblPr"))
        for attr in ("firstRow", "bandRow", "firstCol"):
            tblPr.set(attr, "0")
        for i, w in enumerate(col_w):
            tbl.columns[i].width = Inches(w)
        for r in tbl.rows:
            r.height = Inches(row_h)

        for ri, row in enumerate(rows):
            for ci, val in enumerate(row):
                cell = tbl.cell(ri, ci)
                status = STATUS_WORD.get(val.strip())
                cell.fill.solid()
                cell.fill.fore_color.rgb = rgb(WHITE)
                cell.margin_left = cell.margin_right = Inches(0.08)
                cell.margin_top = cell.margin_bottom = Inches(0.04)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                tf = cell.text_frame
                tf.word_wrap = True
                p = tf.paragraphs[0]
                r = p.add_run()
                r.text = val
                r.font.name = BODY_FONT
                if ri == 0:
                    r.font.size = Pt(11)
                    r.font.bold = True
                    r.font.color.rgb = rgb(AMBER)
                    r.text = val.upper()
                elif status:
                    r.text, col = status
                    r.font.size = Pt(body_size)
                    r.font.bold = True
                    r.font.color.rgb = rgb(col)
                    p.alignment = PP_ALIGN.CENTER
                else:
                    r.font.size = Pt(body_size)
                    r.font.bold = ci == 0
                    r.font.color.rgb = rgb(INK if ci == 0 else BODY)
                self._cell_border(cell, BORDER)
        return gf

    @staticmethod
    def _cell_border(cell, color):
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        for tag in ("a:lnB", "a:lnT"):
            for old in tcPr.findall(qn(tag)):
                tcPr.remove(old)
        from pptx.oxml import parse_xml
        from pptx.oxml.ns import nsdecls
        ln = parse_xml(
            '<a:lnB %s w="12700" cap="flat" cmpd="sng" algn="ctr">'
            '<a:solidFill><a:srgbClr val="%s"/></a:solidFill>'
            '<a:prstDash val="solid"/></a:lnB>' % (nsdecls("a"), color))
        tcPr.append(ln)

    def footer(self, slide, left, right=None):
        self.text(slide, M, 7.10, CW * 0.7, 0.26, left, size=10, color=MUTED)
        if right:
            self.text(slide, M, 7.10, CW, 0.26, right, size=10, color=MUTED,
                      align=PP_ALIGN.RIGHT)


# ----------------------------------------------------------------------------
# Deck assembly
# ----------------------------------------------------------------------------
def fit(text, limit):
    """Trim to `limit` characters, preferring a sentence boundary."""
    text = " ".join(text.split())
    if len(text) <= limit:
        return text
    head = text[:limit]
    cut = max(head.rfind(". "), head.rfind("? "), head.rfind("! "))
    if cut >= limit * 0.55:
        return head[: cut + 1]
    return head.rsplit(" ", 1)[0].rstrip(",;:—-") + "…"


def build(md_path, out_path, log_count=3):
    raw = Path(md_path).read_text(encoding="utf-8")
    meta, decks, body = parse_front_matter(raw)
    body = re.sub(r"<!--.*?-->", "", body, flags=re.S)
    sec = split_sections(body)

    project = meta.get("project", "Project")
    version = meta.get("version", "")
    status = meta.get("status", "")
    updated = meta.get("last_updated", date.today().isoformat())
    owner = meta.get("owner", "")
    started = meta.get("started", "")
    appliance = meta.get("appliance", "")

    d = Deck()
    foot = f"{project}  ·  {version}  ·  as of {updated}"

    # --- 1. Title -----------------------------------------------------
    s = d.slide(dark=True)
    d.text(s, M, 1.75, 9.0, 0.35, f"PROJECT REPORT  ·  {updated}", size=13, bold=True,
           color=AMBER, spc=300, anchor=MSO_ANCHOR.MIDDLE)
    d.text(s, M, 2.25, 11.5, 0.95, project, size=44, bold=True, color=WHITE,
           font=HEAD_FONT, anchor=MSO_ANCHOR.MIDDLE)
    one_liner = ""
    if 1 in sec:
        paras = [p for p in sec[1]["body"].split("\n\n") if p.strip()]
        if paras:
            one_liner = unwrap(strip_md(paras[0]))
    d.text(s, M, 3.20, 11.6, 1.25, fit(one_liner, 390), size=14,
           color=ON_DARK_SUB, line=21, anchor=MSO_ANCHOR.TOP)
    d.text(s, M, 4.55, 10.5, 0.40, f"{version} · {status}".strip(" ·"), size=13,
           color=ON_DARK_MUTED, anchor=MSO_ANCHOR.MIDDLE)
    d.swatches(s)
    d.text(s, M, 5.60, 6.0, 0.32, owner, size=14, bold=True, color=WHITE,
           anchor=MSO_ANCHOR.MIDDLE)
    d.text(s, M, 5.95, 10.0, 0.32, appliance, size=13, color=ON_DARK_MUTED,
           anchor=MSO_ANCHOR.MIDDLE)

    # --- 2. Status at a glance ---------------------------------------
    if 6 in sec:
        s = d.slide()
        d.heading(s, "Where the project stands", f"As of {updated}.")
        stats = [("VERSION", version or "—"), ("STATUS", fit(status, 46) or "—"),
                 ("STARTED", started or "—"), ("LAST UPDATE", updated)]
        cw = (CW - 3 * 0.18) / 4
        for i, (k, v) in enumerate(stats):
            l = M + i * (cw + 0.18)
            d.rect(s, l, 1.62, cw, 0.92, fill=WHITE, line=BORDER)
            d.rect(s, l, 1.62, cw, 0.06, fill=ACCENTS[i % 3])
            d.text(s, l + 0.16, 1.82, cw - 0.32, 0.22, k, size=9.5, bold=True,
                   color=AMBER, spc=200)
            d.text(s, l + 0.16, 2.06, cw - 0.32, 0.42, v, size=12.5, bold=True,
                   color=INK, font=HEAD_FONT, line=15)
        blocks = parse_bold_blocks(sec[6]["body"])
        blocks = [b for b in blocks if b[1]]
        t = 2.76
        per = min(1.12, (6.55 - t) / max(1, len(blocks)))
        for i, (label, txt) in enumerate(blocks):
            top = t + i * (per + 0.14)
            d.rect(s, M, top, CW, per, fill=WHITE, line=BORDER)
            d.rect(s, M, top, 0.055, per, fill=ACCENTS[i % 3])
            d.text(s, M + 0.30, top + 0.14, 2.3, 0.26, label.upper(), size=10.5,
                   bold=True, color=ACCENTS[i % 3], spc=150)
            d.text(s, M + 0.30, top + 0.42, CW - 0.62, per - 0.54,
                   fit(txt, 420), size=12, color=BODY, line=17)
        d.footer(s, foot)

    # --- 3. Why this exists ------------------------------------------
    if 2 in sec:
        qs = parse_numbered(sec[2]["body"])
        blocks = dict(parse_bold_blocks(sec[2]["body"]))
        s = d.slide()
        d.heading(s, "Why this project exists",
                  "The artifact is an instrument. These are what it measures.")
        t = 1.62
        n = max(1, len(qs))
        res = blocks.get("What a result looks like")
        bottom = 6.02 if res else 6.90
        h = min(1.05, (bottom - t - 0.14 * (n - 1)) / n)
        for i, (num, txt) in enumerate(qs):
            d.numbered_card(s, num, t + i * (h + 0.14), h, None,
                            fit(txt, 330), ACCENTS[i % 3])
        if res:
            d.callout(s, t + n * (h + 0.14) + 0.04,
                      "What a result looks like", fit(res, 300), h=0.82)
        d.footer(s, foot)

    # --- 4. Scope -----------------------------------------------------
    if 3 in sec:
        tables = parse_tables(sec[3]["body"])
        if tables:
            rows = tables[0]
            s = d.slide()
            d.heading(s, sec[3]["title"],
                      "Stated explicitly, because the artifact is easy to mistake for the goal.")
            t = 1.70
            h = min(1.00, (6.70 - t - 0.14 * (len(rows) - 1)) / max(1, len(rows)))
            for i, row in enumerate(rows):
                label = row[0] or "—"
                val = row[1] if len(row) > 1 else ""
                top = t + i * (h + 0.14)
                d.rect(s, M, top, CW, h, fill=WHITE, line=BORDER)
                d.rect(s, M, top, 2.45, h, fill=PANEL)
                d.text(s, M + 0.22, top, 2.05, h, label, size=13, bold=True,
                       color=INK, font=HEAD_FONT, anchor=MSO_ANCHOR.MIDDLE)
                d.text(s, M + 2.70, top + 0.10, CW - 2.95, h - 0.20, fit(val, 320),
                       size=12, color=BODY, line=17, anchor=MSO_ANCHOR.MIDDLE)
            d.footer(s, foot)

    # --- 5. Architecture ---------------------------------------------
    if 4 in sec:
        tables = parse_tables(sec[4]["body"])
        blocks = dict(parse_bold_blocks(sec[4]["body"]))
        loop = blocks.get("The loop")
        s = d.slide()
        d.heading(s, sec[4]["title"], "Hardware to public surface, and the gate in between.")
        if tables:
            avail = (4.30 if loop else 5.25)
            d.table(s, 1.62, tables[0], col_w=[1.6, 3.1, 7.4], max_h=avail,
                    body_size=10.5)
        if loop:
            d.callout(s, 5.96, "The loop", fit(loop, 380), h=1.10, size=11.5)
        d.footer(s, foot)

    # --- 6. Related patterns -----------------------------------------
    if 5 in sec:
        tables = parse_tables(sec[5]["body"])
        if tables:
            s = d.slide()
            d.heading(s, "Where this connects to the pattern library",
                      "This project is the live case for work already written up.")
            d.table(s, 1.70, tables[0], col_w=[4.3, 7.8], max_h=4.4, body_size=11)
            d.callout(s, 6.22, "Companion decks",
                      "  ·  ".join(dk.get("title", "") for dk in decks) or "—",
                      h=0.80)
            d.footer(s, foot)

    # --- 7. Milestones ------------------------------------------------
    if 7 in sec:
        tables = parse_tables(sec[7]["body"])
        if tables:
            s = d.slide()
            d.heading(s, "Milestones", "Shipped, in flight, and not yet attempted.")
            d.table(s, 1.70, tables[0], col_w=[1.7, 8.6, 1.8], max_h=5.0,
                    body_size=11)
            d.footer(s, foot)

    # --- 8. Lessons ---------------------------------------------------
    if 8 in sec:
        tables = parse_tables(sec[8]["body"])
        if tables:
            rows = tables[0][1:]
            s = d.slide()
            d.heading(s, "What I've learned",
                      "Promoted out of the daily log. Each one transfers to a project unlike this one.")
            t = 1.60
            n = max(1, len(rows))
            h = min(1.08, (6.95 - t - 0.13 * (n - 1)) / n)
            for i, row in enumerate(rows):
                d.numbered_card(s, row[0], t + i * (h + 0.13), h, None,
                                fit(row[1], 330), ACCENTS[i % 3])
            d.footer(s, foot)

    # --- 9. Open questions --------------------------------------------
    if 9 in sec:
        tables = parse_tables(sec[9]["body"])
        if tables:
            s = d.slide()
            d.heading(s, sec[9]["title"],
                      "Stated plainly, so the pattern is adopted for the right reasons.")
            d.table(s, 1.70, tables[0], col_w=[3.2, 1.5, 7.4], max_h=5.0,
                    body_size=11)
            d.footer(s, foot)

    # --- 10. Recent log ----------------------------------------------
    if 10 in sec:
        entries = parse_log(sec[10]["body"], limit=log_count)
        if entries:
            d.divider("THE LOG", "Recent progress",
                      f"The {len(entries)} most recent updates, newest first.")
        for e in entries:
            s = d.slide()
            d.text(s, M, 0.38, CW, 0.30, e["date"], size=13, bold=True, color=AMBER,
                   spc=250, anchor=MSO_ANCHOR.MIDDLE)
            d.text(s, M, 0.72, CW, 0.62, e["headline"], size=28, bold=True,
                   color=INK, font=HEAD_FONT, anchor=MSO_ANCHOR.MIDDLE)
            fields = [f for f in e["fields"] if f[1]]
            t = 1.58
            n = max(1, len(fields))
            h = min(1.18, (6.95 - t - 0.13 * (n - 1)) / n)
            for i, (label, txt) in enumerate(fields):
                top = t + i * (h + 0.13)
                hi = label.lower().startswith("lesson")
                d.rect(s, M, top, CW, h, fill=PANEL if hi else WHITE, line=BORDER)
                d.rect(s, M, top, 0.055, h, fill=ACCENTS[i % 3])
                d.text(s, M + 0.30, top + 0.13, 2.6, 0.25, label.upper(), size=10.5,
                       bold=True, color=ACCENTS[i % 3], spc=150)
                d.text(s, M + 0.30, top + 0.40, CW - 0.62, h - 0.50, fit(txt, 420),
                       size=11, color=BODY, line=15)
            d.footer(s, foot)

    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    d.prs.save(out_path)
    return out_path, len(d.prs.slides.__iter__.__self__._sldIdLst)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("markdown")
    ap.add_argument("-o", "--out")
    ap.add_argument("-n", "--log-entries", type=int, default=3)
    a = ap.parse_args()
    md = Path(a.markdown)
    if not md.exists():
        sys.exit(f"not found: {md}")
    meta, _, _ = parse_front_matter(md.read_text(encoding="utf-8"))
    name = meta.get("project", md.stem)
    upd = meta.get("last_updated", date.today().isoformat())
    out = a.out or md.with_name(f"Project {name} — Status {upd}.pptx")
    path, n = build(md, out, a.log_entries)
    print(f"wrote {path} ({n} slides)")


if __name__ == "__main__":
    main()
